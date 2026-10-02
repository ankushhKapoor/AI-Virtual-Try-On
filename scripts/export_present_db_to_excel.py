"""
AI Virtual Try-On - DWM Present Database Exporter to Excel and CSV
Extracts live tables from both OLTP (`virtual_tryon`) and DWH (`virtual_tryon_dwh`),
saves them into a multi-sheet, professionally styled Excel (.xlsx) workbook,
and outputs individual CSV files for every table.
"""
import os
import sys
import csv
from datetime import datetime, date
from decimal import Decimal
from typing import Dict, List, Any

# Ensure project root is in sys.path
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from sqlalchemy import text, inspect
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from dwm.connection import oltp_engine, dwh_engine

OUTPUT_DIR = os.path.join(PROJECT_ROOT, "dwm_exports", "present_data")
CSV_DIR = os.path.join(OUTPUT_DIR, "csv")


def sanitize_value_for_excel(val: Any) -> Any:
    """Converts datatypes to Excel-compatible formats."""
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(val, datetime) else val.strftime("%Y-%m-%d")
    if isinstance(val, Decimal):
        return float(val)
    if isinstance(val, (dict, list)):
        import json
        return json.dumps(val)
    return val


def style_worksheet(ws, title: str, row_count: int, col_count: int):
    """Applies professional enterprise styling to Excel worksheets."""
    ws.views.sheetView[0].showGridLines = True
    
    # Header styling: Dark Navy with white text
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Border
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    # Style header row (Row 1)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Zebra striping for data rows
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)

    for row_idx in range(2, row_count + 2):
        ws.row_dimensions[row_idx].height = 20
        use_zebra = (row_idx % 2 == 0)
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill
            # Align numbers right, text left
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(cell.value, (datetime, date)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit column widths with padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    # Freeze top row
    ws.freeze_panes = "A2"


def export_present_database():
    os.makedirs(CSV_DIR, exist_ok=True)
    print(f"[*] Exporting current database tables to:\n    Excel: {OUTPUT_DIR}\n    CSV:   {CSV_DIR}")

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    overview_ws = wb.create_sheet(title="00_Overview_Summary")
    overview_ws.views.sheetView[0].showGridLines = True

    # Tables to extract from DWH (OLAP)
    dwh_tables = [
        ("fact_tryon_event", "Fact_Tryon_Event"),
        ("dim_user", "Dim_User"),
        ("dim_product", "Dim_Product"),
        ("dim_time", "Dim_Time"),
        ("dim_device", "Dim_Device"),
        ("dim_outcome", "Dim_Outcome"),
        ("agg_tryon_daily", "Agg_Tryon_Daily"),
        ("agg_tryon_monthly", "Agg_Tryon_Monthly"),
        ("mining_association_rules", "Mining_Association_Rules"),
        ("mining_quality_correlations", "Mining_Correlations"),
        ("etl_watermark", "ETL_Watermark"),
    ]

    # Tables to extract from OLTP
    oltp_tables = [
        ("vton_jobs", "OLTP_VTON_Jobs"),
        ("users", "OLTP_Users"),
        ("products", "OLTP_Products"),
        ("admins", "OLTP_Admins"),
    ]

    table_summaries = []

    # 1. Export DWH Tables
    with dwh_engine.connect() as conn:
        for tbl_name, sheet_name in dwh_tables:
            try:
                result = conn.execute(text(f"SELECT * FROM `{tbl_name}`"))
                columns = list(result.keys())
                rows = result.fetchall()
            except Exception as e:
                print(f"[!] Warning reading DWH `{tbl_name}`: {e}")
                columns = ["Error"]
                rows = [(str(e),)]

            table_summaries.append(("DWH (OLAP)", tbl_name, sheet_name, len(rows), len(columns)))

            # Write to CSV
            csv_path = os.path.join(CSV_DIR, f"dwh_{tbl_name}.csv")
            with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(columns)
                for r in rows:
                    writer.writerow([sanitize_value_for_excel(v) for v in r])

            # Write to Excel Worksheet
            ws = wb.create_sheet(title=sheet_name)
            ws.append(columns)
            for r in rows:
                ws.append([sanitize_value_for_excel(v) for v in r])
            style_worksheet(ws, sheet_name, len(rows), len(columns))
            print(f"  [DWH] Exported {tbl_name} ({len(rows)} rows) -> CSV & Excel sheet '{sheet_name}'")

    # 2. Export OLTP Tables
    with oltp_engine.connect() as conn:
        for tbl_name, sheet_name in oltp_tables:
            try:
                result = conn.execute(text(f"SELECT * FROM `{tbl_name}`"))
                columns = list(result.keys())
                rows = result.fetchall()
            except Exception as e:
                print(f"[!] Warning reading OLTP `{tbl_name}`: {e}")
                columns = ["Error"]
                rows = [(str(e),)]

            table_summaries.append(("OLTP (Operational)", tbl_name, sheet_name, len(rows), len(columns)))

            # Write to CSV
            csv_path = os.path.join(CSV_DIR, f"oltp_{tbl_name}.csv")
            with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(columns)
                for r in rows:
                    writer.writerow([sanitize_value_for_excel(v) for v in r])

            # Write to Excel Worksheet
            ws = wb.create_sheet(title=sheet_name)
            ws.append(columns)
            for r in rows:
                ws.append([sanitize_value_for_excel(v) for v in r])
            style_worksheet(ws, sheet_name, len(rows), len(columns))
            print(f"  [OLTP] Exported {tbl_name} ({len(rows)} rows) -> CSV & Excel sheet '{sheet_name}'")

    # 3. Create Overview Sheet
    overview_headers = ["Database Tier", "Table Name", "Excel Sheet Name", "Total Record Count", "Total Columns"]
    overview_ws.append(["AI Virtual Try-On - DWM Data Warehouse & OLTP Database Summary"])
    overview_ws.append([f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    overview_ws.append([])
    overview_ws.append(overview_headers)

    for item in table_summaries:
        overview_ws.append(list(item))

    # Style Overview Sheet
    overview_ws.merge_cells("A1:E1")
    title_cell = overview_ws["A1"]
    title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="1F497D")
    overview_ws.row_dimensions[1].height = 32

    overview_ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="555555")

    # Style table headers in row 4
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, 6):
        cell = overview_ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    overview_ws.row_dimensions[4].height = 25

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    for row_idx in range(5, 5 + len(table_summaries)):
        overview_ws.row_dimensions[row_idx].height = 20
        use_zebra = (row_idx % 2 == 0)
        for col_idx in range(1, 6):
            cell = overview_ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if use_zebra:
                cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
            if col_idx in (4, 5):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in overview_ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2:
                max_len = max(max_len, len(str(cell.value or "")))
        overview_ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    excel_file_path = os.path.join(OUTPUT_DIR, "Virtual_TryOn_DWM_Present_Data.xlsx")
    wb.save(excel_file_path)
    print(f"\n[SUCCESS] Master Excel Workbook created: {excel_file_path}")
    print(f"[SUCCESS] CSV files exported to: {CSV_DIR}\n")


if __name__ == "__main__":
    export_present_database()
